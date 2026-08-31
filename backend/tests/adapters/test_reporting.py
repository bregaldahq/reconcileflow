from __future__ import annotations

import io
from datetime import UTC, date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from reconcileflow.adapters.reporting import XlsxReportBuilder
from reconcileflow.application.models import ManualReviewEvent, ReportPayload
from reconcileflow.domain.matching import ExactCompositeRule, ReconciliationEngine
from reconcileflow.domain.models import (
    CanonicalTransaction,
    DatasetFingerprint,
    DatasetSide,
    ManualAction,
    Money,
    ReconciliationRun,
    RuleReference,
    RunStatus,
)


def test_xlsx_report_has_required_sheets_precision_and_deterministic_bytes() -> None:
    created = datetime(2026, 1, 1, tzinfo=UTC)
    source = CanonicalTransaction(
        "a:1",
        "a",
        DatasetSide.A,
        Money(Decimal("1234.56"), "BRL"),
        date(2026, 1, 1),
        "1",
        "DOC",
        "REF",
    )
    target = CanonicalTransaction(
        "b:1",
        "b",
        DatasetSide.B,
        Money(Decimal("1234.56"), "BRL"),
        date(2026, 1, 1),
        "1",
        "DOC",
        "REF",
    )
    run = ReconciliationRun(
        "run-1",
        DatasetFingerprint.from_bytes(b"a"),
        DatasetFingerprint.from_bytes(b"b"),
        RuleReference("standard", "1"),
        "br-v1",
        RunStatus.COMPLETED,
        created,
        created,
        "corr",
        "idem",
    )
    result = ReconciliationEngine((ExactCompositeRule(),)).reconcile(
        run_id="run-1", transactions=(source, target)
    )
    review = ManualReviewEvent(
        "review-1",
        "run-1",
        result.decisions[0].decision_id,
        ManualAction.APPROVE,
        "analyst",
        "=formula-like reason",
        "review-idem",
        created,
    )
    payload = ReportPayload(run, result.decisions, (source, target), (review,))
    builder = XlsxReportBuilder()
    first = builder.build(payload)
    second = builder.build(payload)
    assert first == second

    workbook = load_workbook(io.BytesIO(first), read_only=False, data_only=False)
    assert workbook.sheetnames == [
        "Resumo",
        "Conciliados",
        "Sem correspondência",
        "Mais de uma opção",
        "Duplicidades",
        "Revisão manual",
    ]
    matched = workbook["Conciliados"]
    assert matched["A1"].value == "ID da decisão"
    assert matched["B2"].value == "Conciliado"
    assert matched["H2"].value == "Todos os campos canônicos configurados são iguais."
    assert matched["J2"].value == 1234.56
    assert matched["J2"].number_format == "#,##0.00;[Red]-#,##0.00"
    manual = workbook["Revisão manual"]
    assert manual["B2"].value == "Aprovado por analyst"
    assert manual["H2"].value == "'=formula-like reason"
    workbook.close()
