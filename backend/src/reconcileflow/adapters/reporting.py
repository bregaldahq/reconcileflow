"""Professional deterministic XLSX reconciliation report adapter."""

from __future__ import annotations

import io
import re
import zipfile
from collections import Counter
from datetime import UTC
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from reconcileflow.application.models import ReportPayload
from reconcileflow.domain.models import MatchStatus, ReconciliationDecision

_HEADER_FILL = PatternFill("solid", fgColor="173F46")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_MONEY_FORMAT = "#,##0.00;[Red]-#,##0.00"
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_MATCH_STATUS_LABELS = {
    MatchStatus.MATCHED: "Conciliado",
    MatchStatus.SPLIT_MATCHED: "Conciliado em grupo",
    MatchStatus.UNMATCHED: "Sem correspondência",
    MatchStatus.AMBIGUOUS: "Mais de uma possibilidade",
    MatchStatus.DUPLICATE: "Possível duplicidade",
    MatchStatus.MANUAL_REVIEW: "Revisão necessária",
}
_RUN_STATUS_LABELS = {
    "PENDING": "Pendente",
    "QUEUED": "Na fila",
    "RUNNING": "Em processamento",
    "COMPLETED": "Concluída",
    "FAILED": "Falhou",
}
_EXPLANATION_TRANSLATIONS = {
    "All configured canonical fields are equal.": (
        "Todos os campos canônicos configurados são iguais."
    ),
    "Multiple equally valid candidates require manual review.": (
        "Há mais de uma possibilidade igualmente válida; a decisão precisa de revisão."
    ),
    "Amount and date are within configured tolerances.": (
        "O valor e a data estão dentro dos limites configurados."
    ),
    "A bounded group conserves the reconciled amount.": (
        "Um grupo dentro dos limites conserva o valor conciliado."
    ),
    "Canonical records are duplicated on the same dataset side.": (
        "Existem registros repetidos no mesmo lado da conciliação."
    ),
    "Candidate search reached its configured safety limit.": (
        "A busca de possibilidades atingiu o limite de segurança configurado."
    ),
    "No configured automatic rule produced a candidate.": (
        "Nenhuma regra automática configurada encontrou uma possibilidade."
    ),
}
_FIELD_LABELS = {
    "amount": "valor",
    "currency": "moeda",
    "date": "data",
    "document": "documento",
    "reference": "referência",
}
_RULE_LABELS = {
    "standard": "Conciliação explicável padrão",
    "exact_composite": "Correspondência exata",
    "amount_date_tolerance": "Tolerância de valor e data",
    "bounded_split": "Conciliação em grupo",
}


def _safe_text(value: str) -> str:
    if value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _translate_explanation(value: str) -> str:
    return _EXPLANATION_TRANSLATIONS.get(value, value)


def _rule_label(value: str) -> str:
    return _RULE_LABELS.get(value, "Regra personalizada")


def _translate_evidence(value: str) -> str:
    patterns = (
        (
            r"(?:candidate sources|possibilidade origens)=(.+) "
            r"(?:targets|destinos)=(.+)",
            "possibilidade: origem {} → destino {}",
        ),
        (r"amount delta (\S+) <= (\S+)", "diferença de valor {} ≤ {}"),
        (r"date (?:delta|span) (\S+) <= (\S+)", "diferença de data {} ≤ {}"),
        (r"group size (\S+) <= (\S+)", "tamanho do grupo {} ≤ {}"),
        (r"duplicate count (\d+)", "quantidade de duplicados: {}"),
    )
    for pattern, template in patterns:
        matched = re.fullmatch(pattern, value)
        if matched:
            return template.format(*matched.groups())
    equal = re.fullmatch(r"([a-z_]+): (?:equal|igual)", value)
    if equal:
        return f"{_FIELD_LABELS.get(equal.group(1), equal.group(1))}: igual"
    literal_translations = {
        "bounded search limit reached": "limite da busca restrita atingido",
        "matching pipeline exhausted": "todas as regras configuradas foram avaliadas",
    }
    return literal_translations.get(value, value.replace(" <= ", " ≤ "))


def _prepare_sheet(sheet: Worksheet, headers: tuple[str, ...]) -> None:
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).coordinate}"
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24


def _finish_sheet(sheet: Worksheet) -> None:
    widths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for column_index, cell in enumerate(row, start=1):
            rendered = "" if cell.value is None else str(cell.value)
            widths[column_index] = min(max(widths.get(column_index, 0), len(rendered) + 2), 48)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = max(width, 12)


class XlsxReportBuilder:
    def build(self, payload: ReportPayload) -> bytes:
        workbook = Workbook()
        active_sheet = workbook.active
        if active_sheet is not None:
            workbook.remove(active_sheet)
        created = payload.run.created_at.astimezone(UTC).replace(tzinfo=None)
        workbook.properties.created = created
        workbook.properties.modified = created
        workbook.properties.creator = "ReconcileFlow"
        workbook.properties.title = f"Conciliação {payload.run.run_id}"

        self._summary(workbook, payload)
        partitions = {
            "Conciliados": {MatchStatus.MATCHED, MatchStatus.SPLIT_MATCHED},
            "Sem correspondência": {MatchStatus.UNMATCHED},
            "Mais de uma opção": {MatchStatus.AMBIGUOUS},
            "Duplicidades": {MatchStatus.DUPLICATE},
            "Revisão manual": {MatchStatus.MANUAL_REVIEW},
        }
        transactions = {item.transaction_id: item for item in payload.transactions}
        for title, statuses in partitions.items():
            sheet = workbook.create_sheet(title)
            headers = (
                "ID da decisão",
                "Resultado",
                "IDs da origem",
                "IDs do destino",
                "Regra aplicada",
                "Versão da regra",
                "Confiança",
                "Explicação",
                "Evidências",
                "Valor da origem",
                "Valor do destino",
            )
            _prepare_sheet(sheet, headers)
            for decision in payload.decisions:
                if decision.status not in statuses:
                    continue
                source_amount = sum(
                    (transactions[item].money.amount for item in decision.source_ids),
                    start=Decimal("0"),
                )
                target_amount = sum(
                    (transactions[item].money.amount for item in decision.target_ids),
                    start=Decimal("0"),
                )
                sheet.append(self._decision_row(decision, source_amount, target_amount))
                sheet.cell(sheet.max_row, 10).number_format = _MONEY_FORMAT
                sheet.cell(sheet.max_row, 11).number_format = _MONEY_FORMAT
            if title == "Revisão manual":
                for review in payload.reviews:
                    sheet.append(
                        (
                            review.decision_id,
                            (
                                "Aprovado por "
                                if review.action.value == "APPROVE"
                                else "Recusado por "
                            )
                            + _safe_text(review.actor_id),
                            "",
                            "",
                            "decisão manual",
                            "1",
                            "",
                            _safe_text(review.reason),
                            review.created_at.isoformat(),
                            None,
                            None,
                        )
                    )
            _finish_sheet(sheet)

        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        return self._canonical_zip(buffer.getvalue())

    def _summary(self, workbook: Workbook, payload: ReportPayload) -> None:
        sheet = workbook.create_sheet("Resumo")
        _prepare_sheet(sheet, ("Métrica", "Valor"))
        counts = Counter(decision.status for decision in payload.decisions)
        rows = (
            ("ID da execução", payload.run.run_id),
            ("Status", _RUN_STATUS_LABELS.get(payload.run.status.value, payload.run.status.value)),
            ("Impressão digital da origem (A)", payload.run.dataset_a.sha256),
            ("Impressão digital do destino (B)", payload.run.dataset_b.sha256),
            (
                "Conjunto de regras",
                f"{_rule_label(payload.run.rule_set.rule_id)} · versão "
                f"{payload.run.rule_set.version}",
            ),
            ("Normalização", payload.run.normalization_version),
            *((_MATCH_STATUS_LABELS[status], counts[status]) for status in MatchStatus),
            ("Decisões manuais", len(payload.reviews)),
        )
        for metric, value in rows:
            sheet.append((_safe_text(str(metric)), _safe_text(str(value))))
        _finish_sheet(sheet)

    @staticmethod
    def _decision_row(
        decision: ReconciliationDecision,
        source_amount: Decimal,
        target_amount: Decimal,
    ) -> tuple[object, ...]:
        return (
            decision.decision_id,
            _MATCH_STATUS_LABELS[decision.status],
            ", ".join(decision.source_ids),
            ", ".join(decision.target_ids),
            _rule_label(decision.rule.rule_id) if decision.rule else "",
            decision.rule.version if decision.rule else "",
            str(decision.confidence) if decision.confidence is not None else "",
            _safe_text(_translate_explanation(decision.explanation.summary)),
            " | ".join(
                _safe_text(_translate_evidence(item)) for item in decision.explanation.evidence
            ),
            source_amount,
            target_amount,
        )

    @staticmethod
    def _canonical_zip(content: bytes) -> bytes:
        source = io.BytesIO(content)
        target = io.BytesIO()
        with (
            zipfile.ZipFile(source, "r") as input_zip,
            zipfile.ZipFile(
                target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
            ) as output_zip,
        ):
            for name in sorted(input_zip.namelist()):
                info = zipfile.ZipInfo(name, _FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                output_zip.writestr(info, input_zip.read(name))
        return target.getvalue()
