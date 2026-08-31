"""Bounded CSV and XLSX parsing at the infrastructure edge."""

from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import PurePath
from typing import cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from reconcileflow.application.errors import DatasetValidationError, ValidationIssue
from reconcileflow.application.models import ParsedDataset
from reconcileflow.domain.models import DatasetFingerprint


@dataclass(frozen=True, slots=True)
class IngestionLimits:
    max_bytes: int = 10 * 1024 * 1024
    max_uncompressed_bytes: int = 50 * 1024 * 1024
    max_rows: int = 100_000
    max_columns: int = 200
    max_cell_characters: int = 20_000


class TabularDatasetParser:
    def __init__(self, limits: IngestionLimits | None = None) -> None:
        self.limits = limits or IngestionLimits()

    def parse(self, *, filename: str, content_type: str, content: bytes) -> ParsedDataset:
        if not content:
            self._fail("empty_file", "o arquivo enviado está vazio")
        if len(content) > self.limits.max_bytes:
            self._fail("file_too_large", f"o arquivo ultrapassa {self.limits.max_bytes} bytes")
        suffix = PurePath(filename).suffix.lower()
        if suffix == ".csv":
            if content_type not in {"text/csv", "application/csv", "application/vnd.ms-excel"}:
                self._fail(
                    "invalid_content_type", "o arquivo CSV possui um tipo de mídia não aceito"
                )
            headers, rows = self._parse_csv(content)
        elif suffix == ".xlsx":
            if content_type not in {
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "application/octet-stream",
            }:
                self._fail(
                    "invalid_content_type", "o arquivo XLSX possui um tipo de mídia não aceito"
                )
            headers, rows = self._parse_xlsx(content)
        else:
            self._fail("unsupported_file_type", "somente arquivos .csv e .xlsx são aceitos")
        return ParsedDataset(
            fingerprint=DatasetFingerprint.from_bytes(content),
            headers=headers,
            rows=rows,
            byte_size=len(content),
        )

    def _parse_csv(self, content: bytes) -> tuple[tuple[str, ...], tuple[dict[str, object], ...]]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise DatasetValidationError(
                (
                    ValidationIssue(
                        "invalid_encoding", "o arquivo CSV deve usar a codificação UTF-8"
                    ),
                )
            ) from error
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text, newline=""), dialect)
        try:
            raw_headers = next(reader)
        except StopIteration:
            self._fail("missing_header", "o arquivo CSV não possui uma linha de cabeçalho")
        headers = self._validate_headers(tuple(raw_headers))
        rows: list[dict[str, object]] = []
        for row_number, values in enumerate(reader, start=2):
            if row_number - 1 > self.limits.max_rows:
                self._fail("too_many_rows", f"a fonte ultrapassa {self.limits.max_rows} linhas")
            if len(values) != len(headers):
                self._fail(
                    "invalid_column_count",
                    (
                        f"a linha {row_number} possui {len(values)} valores; "
                        f"eram esperados {len(headers)}"
                    ),
                    row=row_number,
                )
            rows.append(
                {
                    header: self._bounded_cell(value, row_number, header)
                    for header, value in zip(headers, values, strict=True)
                }
            )
        return headers, tuple(rows)

    def _parse_xlsx(self, content: bytes) -> tuple[tuple[str, ...], tuple[dict[str, object], ...]]:
        if not content.startswith(b"PK"):
            self._fail("invalid_signature", "a assinatura do arquivo XLSX é inválida")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                expanded = sum(item.file_size for item in archive.infolist())
                if expanded > self.limits.max_uncompressed_bytes:
                    self._fail(
                        "expanded_file_too_large",
                        "o conteúdo expandido do XLSX ultrapassa o limite configurado",
                    )
                if any(
                    item.filename.lower().endswith("vbaproject.bin") for item in archive.infolist()
                ):
                    self._fail("macros_not_allowed", "planilhas com macros não são aceitas")
        except zipfile.BadZipFile as error:
            raise DatasetValidationError(
                (ValidationIssue("invalid_xlsx", "o pacote do arquivo XLSX é inválido"),)
            ) from error
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        except (OSError, ValueError, KeyError) as error:
            raise DatasetValidationError(
                (ValidationIssue("invalid_xlsx", "não foi possível interpretar a planilha"),)
            ) from error
        try:
            active_sheet = workbook.active
            if active_sheet is None:
                self._fail("missing_sheet", "o arquivo XLSX não possui uma planilha ativa")
            worksheet = cast(Worksheet, active_sheet)
            iterator = worksheet.iter_rows(values_only=True)
            try:
                raw_headers = next(iterator)
            except StopIteration:
                self._fail("missing_header", "o arquivo XLSX não possui uma linha de cabeçalho")
            headers = self._validate_headers(
                tuple("" if value is None else str(value) for value in raw_headers)
            )
            rows: list[dict[str, object]] = []
            for row_number, values in enumerate(iterator, start=2):
                if row_number - 1 > self.limits.max_rows:
                    self._fail("too_many_rows", f"a fonte ultrapassa {self.limits.max_rows} linhas")
                padded = (*values, *(None for _ in range(max(0, len(headers) - len(values)))))
                selected = padded[: len(headers)]
                rows.append(
                    {
                        header: self._bounded_cell(self._safe_xlsx_value(value), row_number, header)
                        for header, value in zip(headers, selected, strict=True)
                    }
                )
            return headers, tuple(rows)
        finally:
            workbook.close()

    def _validate_headers(self, headers: tuple[str, ...]) -> tuple[str, ...]:
        cleaned = tuple(header.strip() for header in headers)
        if not cleaned or any(not header for header in cleaned):
            self._fail("invalid_header", "os cabeçalhos não podem ficar vazios")
        if len(cleaned) > self.limits.max_columns:
            self._fail("too_many_columns", f"a fonte ultrapassa {self.limits.max_columns} colunas")
        if len(cleaned) != len(set(cleaned)):
            self._fail("duplicate_header", "os cabeçalhos devem ser únicos após remover espaços")
        return cleaned

    def _safe_xlsx_value(self, value: object) -> object:
        if isinstance(value, float):
            return format(value, ".15g")
        if isinstance(value, Decimal):
            return str(value)
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return value

    def _bounded_cell(self, value: object, row: int, field: str) -> object:
        if isinstance(value, str) and len(value) > self.limits.max_cell_characters:
            self._fail(
                "cell_too_large",
                f"a célula ultrapassa {self.limits.max_cell_characters} caracteres",
                row=row,
                field=field,
            )
        return value

    @staticmethod
    def _fail(code: str, message: str, *, row: int | None = None, field: str | None = None) -> None:
        raise DatasetValidationError((ValidationIssue(code, message, row, field),))
